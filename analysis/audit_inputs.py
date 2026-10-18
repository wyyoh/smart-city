from __future__ import annotations

import collections
import hashlib
import importlib.util
import json
import math
from pathlib import Path
import re
import unicodedata

import openpyxl
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'analysis'
OUT.mkdir(exist_ok=True)


def norm(value):
    if value is None:
        return None
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(value)))


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


source = next(ROOT.glob('*.xlsx'))
book = openpyxl.load_workbook(source, data_only=True)
formula_book = openpyxl.load_workbook(source, data_only=False)
summary = {'source': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'sheets': {}}
raw = {}
for sheet in book:
    grid = [list(row) for row in sheet.values]
    expanded = [row[:] for row in grid]
    for region in sheet.merged_cells.ranges:
        value = grid[region.min_row - 1][region.min_col - 1]
        for row in range(region.min_row - 1, region.max_row):
            for col in range(region.min_col - 1, region.max_col):
                expanded[row][col] = value
    errors = [{'cell': cell.coordinate, 'value': cell.value, 'formula': formula_book[sheet.title][cell.coordinate].value}
              for row in sheet for cell in row if cell.data_type == 'e']
    raw[sheet.title] = {'values': grid, 'expanded': expanded}
    summary['sheets'][sheet.title] = {
        'rows': sheet.max_row, 'columns': sheet.max_column,
        'nonempty_rows_after_headers': sum(any(v is not None for v in row) for row in grid[2:]),
        'formula_errors': len(errors), 'error_examples': errors[:8],
    }

configs = {
    '体积指标': dict(source=0, grade=1, product=2, specimen=3, oil=6, targets={'stability': 16, 'VV': 20}),
    '浸水马歇尔': dict(source=0, grade=1, product=2, specimen=3, oil=4, targets={'residual_ratio_raw': 16}),
    '冻融劈裂': dict(source=4, grade=5, product=6, specimen=7, oil=8, targets={'TSR_raw': 20}),
    '高温车辙': dict(source=4, grade=5, product=6, specimen=7, oil=8, targets={'DS': 10}),
    '低温弯曲': dict(source=0, grade=1, product=2, specimen=3, oil=4, targets={'strain': 7, 'flexural_strength': 6, 'stiffness': 8}),
}
records = {}
for name, cfg in configs.items():
    items = []
    original = raw[name]['values']
    for rowno, row in enumerate(raw[name]['expanded'][2:], 3):
        if row[cfg['source']] is None or row[cfg['specimen']] is None:
            continue
        item = {'sheet': name, 'row': rowno}
        for key in ['source', 'grade', 'product', 'specimen', 'oil']:
            item[key] = row[cfg[key]]
        item['grade_key'] = norm(item['grade'])
        item['product_key'] = norm(item['product'])
        item['targets'] = {key: original[rowno - 1][col] for key, col in cfg['targets'].items()}
        items.append(item)
    records[name] = items
    groups = collections.Counter((norm(x['source']), x['grade_key'], x['product_key'], round(x['oil'], 6) if numeric(x['oil']) else str(x['oil'])) for x in items)
    summary['sheets'][name].update({
        'specimen_records': len(items),
        'sources': dict(collections.Counter(norm(x['source']) for x in items)),
        'distinct_source_grade_product_oil_keys': len(groups),
        'numeric_targets': {key: sum(numeric(x['targets'][key]) for x in items) for key in cfg['targets']},
    })

grade_sheet = book['级配']
blocks = []
for region in grade_sheet.merged_cells.ranges:
    if region.min_col == 2 and region.max_col == 2 and region.min_row >= 3:
        start, end = region.min_row, region.max_row
        rows = raw['级配']['expanded'][start - 1:end]
        first = rows[0]
        components = []
        for rowno, row in enumerate(rows, start):
            if row[2] in ('粗集料', '细集料', '矿粉', '粗RAP料', '细RAP料'):
                components.append({'row': rowno, 'type': row[2], 'size': row[3], 'new_old_fraction': row[4],
                                   'within_fraction': row[5], 'sieving': row[6:18],
                                   'apparent_density': row[18], 'bulk_density': row[19]})
        rap_fracs = {float(x['new_old_fraction']) for x in components if 'RAP' in str(x['type']) and numeric(x['new_old_fraction'])}
        new_fracs = {float(x['new_old_fraction']) for x in components if 'RAP' not in str(x['type']) and numeric(x['new_old_fraction'])}
        blocks.append({'source': norm(first[0]), 'grade': first[1], 'grade_key': norm(first[1]),
                       'row_start': start, 'row_end': end, 'rap_fractions': sorted(rap_fracs),
                       'new_fractions': sorted(new_fracs), 'design_oil': first[28], 'components': components})
blocks.sort(key=lambda x: x['row_start'])
grade_keys = {(x['source'], x['grade_key']) for x in blocks}
summary['gradation'] = {'blocks': len(blocks), 'unique_source_grade_keys': len(grade_keys),
    'source_distribution': dict(collections.Counter(x['source'] for x in blocks)),
    'rap_fraction_distribution': dict(collections.Counter(str(x['rap_fractions']) for x in blocks)),
    'grade_family_distribution': dict(collections.Counter((re.search(r'(AC|SMA|OGFC)[-－]?\d+', x['grade_key']) or ['other'])[0] for x in blocks))}
for name, items in records.items():
    matched = [x for x in items if (norm(x['source']), x['grade_key']) in grade_keys]
    summary['sheets'][name]['exact_source_grade_matched_rows'] = len(matched)
    summary['sheets'][name]['unmatched_keys'] = sorted({(norm(x['source']), x['grade_key']) for x in items if (norm(x['source']), x['grade_key']) not in grade_keys}, key=str)

for filename, data in [('workbook_summary.json', summary), ('workbook_raw.json', raw), ('performance_records.json', records), ('gradation_blocks.json', blocks)]:
    (OUT / filename).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')

pdf_path = next(ROOT.glob('*.pdf'))
reader = PdfReader(pdf_path)
(OUT / 'task_source.txt').write_text('\n\n'.join(f'第{i+1}页\n{p.extract_text()}' for i, p in enumerate(reader.pages)), encoding='utf-8')
if importlib.util.find_spec('pypdfium2'):
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(str(pdf_path))
    for i in range(5):
        document[i].render(scale=1.5).to_pil().save(OUT / f'task_page_{i+1}.png')

print(json.dumps(summary, ensure_ascii=False, indent=2))
print('BLOCK PREVIEW', json.dumps([{k: v for k, v in x.items() if k != 'components'} for x in blocks], ensure_ascii=False))
