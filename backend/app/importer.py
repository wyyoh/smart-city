"""Versioned, explicit-path importer. Original Excel files are never modified."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import json
import math
from pathlib import Path
import re
import unicodedata

import openpyxl
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select

from .domain.contracts import METRICS
from .settings import ROOT
from .storage import Dataset, Measurement, Recipe, SourceRow, engine, initialize

IMPORTER_VERSION = '1.1.1'
SIEVES = [26.5, 19, 16, 13.2, 9.5, 4.75, 2.36, 1.18, .6, .3, .15, .075]
LAYOUT = {
    '体积指标': dict(source=0, grade=1, product=2, specimen=3, oil=6,
                 targets={'stability': 16, 'flow': 17, 'VV': 20, 'VMA': 21}),
    '浸水马歇尔': dict(source=0, grade=1, product=2, specimen=3, oil=4, targets={'residual': 16}),
    '冻融劈裂': dict(source=4, grade=5, product=6, specimen=7, oil=8, targets={'TSR': 20}),
    '高温车辙': dict(source=4, grade=5, product=6, specimen=7, oil=8, targets={'DS': 10}),
    '低温弯曲': dict(source=0, grade=1, product=2, specimen=3, oil=4, targets={'strain': 7}),
}
TARGET_HEADERS = {'stability': '稳定度', 'flow': '流值', 'VV': '空隙率', 'VMA': '矿料间隙率',
                  'residual': '残留稳定度', 'TSR': 'TSR', 'DS': '动稳定度', 'strain': '最大弯拉应变'}


def norm(value):
    return '' if value is None else re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(value)))


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def stable_id(prefix, value):
    return prefix + hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:20]


def expanded_grid(sheet):
    raw = [list(row) for row in sheet.values]
    expanded = [r[:] for r in raw]
    for region in sheet.merged_cells.ranges:
        v = raw[region.min_row - 1][region.min_col - 1]
        for r in range(region.min_row - 1, region.max_row):
            for c in range(region.min_col - 1, region.max_col):
                expanded[r][c] = v
    return raw, expanded


def verify_layout(book):
    required = set(LAYOUT) | {'设计流程', '级配', '集料筛分及密度'}
    if required - set(book.sheetnames):
        raise ValueError(f'工作簿缺少工作表：{sorted(required - set(book.sheetnames))}')
    for sheet, cfg in LAYOUT.items():
        s = book[sheet]
        if norm(s.cell(1, cfg['source'] + 1).value) != '试件来源':
            raise ValueError(f'{sheet}来源列不符合布局v1，禁止猜测导入')
        for metric, col in cfg['targets'].items():
            if TARGET_HEADERS[metric] not in norm(s.cell(2, col + 1).value):
                raise ValueError(f'{sheet}的{metric}表头变化，请更新导入契约')
    if [book['级配'].cell(2, c).value for c in range(7, 19)] != SIEVES:
        raise ValueError('级配筛孔列表不符合契约')


def parse_gradations(book):
    _, expanded = expanded_grid(book['级配'])
    _, raw_materials = expanded_grid(book['集料筛分及密度'])
    material_index = defaultdict(list)
    for n, row in enumerate(raw_materials[2:], 3):
        material_index[(norm(row[0]), norm(row[1]), norm(row[2]))].append((n, row[3:15]))
    blocks = {}
    mismatches = []
    for region in sorted(book['级配'].merged_cells.ranges, key=lambda x: x.min_row):
        if region.min_col != 2 or region.max_col != 2 or region.min_row < 3:
            continue
        start = region.min_row
        rows = expanded[start - 1:region.max_row]
        first = rows[0]
        source, grade = norm(first[0]), norm(first[1])
        components, rap_values, profile = [], set(), []
        flags = []
        for r, row in enumerate(rows[:9], start):
            if row[2] not in ['粗集料', '细集料', '矿粉', '粗RAP料', '细RAP料']:
                continue
            fraction = float(row[4]) * float(row[5]) if numeric(row[4]) and numeric(row[5]) else None
            curve = [float(v) if numeric(v) else None for v in row[6:18]]
            component = {'row': r, 'type': row[2], 'size': row[3], 'group_fraction': row[4],
                         'within_fraction': row[5], 'nominal_weight': fraction, 'curve': curve,
                         'density_apparent': row[18], 'density_bulk': row[19]}
            components.append(component)
            if 'RAP' in str(row[2]) and numeric(row[4]):
                rap_values.add(float(row[4]))
            if fraction and fraction > 0:
                profile.append([norm(row[2]), norm(row[3]), [round(v, 5) if v is not None else None for v in curve]])
                exact = material_index[(source, norm(row[2]), norm(row[3]))]
                if exact and not any(all(numeric(a) and numeric(b) and abs(a-b) < 1e-6 for a,b in zip(curve, rawcurve)) for _,rawcurve in exact):
                    mismatches.append({'source': source, 'gradation': grade, 'grade_row': r,
                                       'raw_rows': [n for n,_ in exact], 'field': 'curve'})
                    flags.append('component_curve_requires_review')
        stored_curve = rows[9][6:18] if len(rows) > 9 else []
        curve = [float(v) if numeric(v) else None for v in stored_curve]
        if not components or not any(c['nominal_weight'] is not None for c in components):
            flags.append('preblended_curve_without_component_weights')
        # A profile includes every non-zero component. A shared filler alone never merges sources.
        complete_profile = bool(profile) and all(all(v is not None for v in item[2]) for item in profile)
        profile_id = stable_id('profile_', sorted(profile, key=str)) if complete_profile else None
        block = {'source': source, 'gradation': grade, 'start_row': start, 'end_row': region.max_row,
                 'components': components, 'curve': curve, 'sieves_mm': SIEVES,
                 'rap_nominal_pct': next(iter(rap_values))*100 if len(rap_values) == 1 else 0 if components and not rap_values else None,
                 'design_oil_stone_ratio_pct': first[28] if numeric(first[28]) else None,
                 'rap_mass_basis': 'unconfirmed', 'profile_id': profile_id,
                 'flags': sorted(set(flags)), 'curve_reference': f'级配!G{start+9}:R{start+9}'}
        if (source, grade) in blocks:
            raise ValueError(f'重复级配键需显式消歧：{source}/{grade}')
        blocks[(source, grade)] = block
    return blocks, mismatches


def operational_product(source, product, specimen, sheet):
    if source == 'B' and '方案' in product:
        return re.sub(r'-\d+$', '', specimen) if sheet == '低温弯曲' else specimen
    return product


def import_workbook(path: Path, session: Session):
    path = path.resolve(strict=True)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    dsid = stable_id('ds_', [digest, IMPORTER_VERSION])
    existing = session.get(Dataset, dsid)
    if existing:
        return existing, False
    book = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    verify_layout(book)
    blocks, mismatches = parse_gradations(book)
    ds = Dataset(id=dsid, filename=path.name, sha256=digest, importer_version=IMPORTER_VERSION)
    session.add(ds)
    session.flush()
    sheets, errors, counts, groups = {}, [], {}, defaultdict(lambda: defaultdict(list))
    for sheet in book:
        raw, expanded = expanded_grid(sheet)
        sheets[sheet.title] = {'rows': sheet.max_row, 'columns': sheet.max_column,
                               'merged_ranges': [str(r) for r in sheet.merged_cells.ranges]}
        for rowno, row in enumerate(raw, 1):
            if not any(v is not None for v in row):
                continue
            cells = {}
            for col, value in enumerate(row, 1):
                formula = formulas[sheet.title].cell(rowno, col)
                if value is not None or formula.data_type == 'f':
                    coordinate = f'{get_column_letter(col)}{rowno}'
                    cells[coordinate] = {'value': value, 'formula': formula.value if formula.data_type == 'f' else None,
                                         'number_format': sheet.cell(rowno, col).number_format,
                                         'error': sheet.cell(rowno, col).data_type == 'e'}
                    if sheet.cell(rowno, col).data_type == 'e':
                        errors.append({'sheet': sheet.title, 'cell': coordinate, 'value': value})
            session.add(SourceRow(id=f'{dsid}:{sheet.title}:{rowno}', dataset_id=dsid, sheet=sheet.title,
                                  row=rowno, payload={'cells': cells}))
        if sheet.title not in LAYOUT:
            continue
        cfg, count = LAYOUT[sheet.title], 0
        for n, row in enumerate(expanded[2:], 3):
            if row[cfg['source']] is None or row[cfg['specimen']] is None:
                continue
            count += 1
            source, grade, product, specimen = [norm(row[cfg[k]]) for k in ['source','grade','product','specimen']]
            product = operational_product(source, product, specimen, sheet.title)
            oil = round(float(row[cfg['oil']]), 6) if numeric(row[cfg['oil']]) else None
            key = (source, grade, product, oil)
            # Retain missing-label groups, using original cells rather than merged target copies.
            groups[key]['_records'].append({'sheet': sheet.title, 'row': n, 'specimen': specimen})
            for metric, col in cfg['targets'].items():
                value = raw[n-1][col]
                if numeric(value):
                    canonical = float(value) * (100 if metric == 'TSR' else 1)
                    groups[key][metric].append({'value': canonical, 'raw_value': value,
                        'sheet': sheet.title, 'row': n, 'cell': f'{get_column_letter(col+1)}{n}',
                        'transform': 'fraction_to_percent_x100' if metric == 'TSR' else 'identity',
                        'specimen': specimen})
        counts[sheet.title] = count
    fixtures = json.loads((ROOT/'configs/cases.json').read_text(encoding='utf-8'))['candidates']
    fixture_index = {(x['source'], norm(x['gradation']), norm(x['product']), round(x['oil_stone_ratio_pct'],6)): x['id'] for x in fixtures}
    suspected, no_grade = [], []
    for key, metric_items in groups.items():
        source, grade, product, oil = key
        recipe_id = stable_id('mix_', [dsid, *key])
        gradation = blocks.get((source, grade))
        flags = list(gradation['flags']) if gradation else ['gradation_key_unmatched']
        if not gradation:
            no_grade.append(list(key))
        if source == 'T' and any(x['sheet'] == '低温弯曲' and 143 <= x['row'] <= 148 for x in metric_items['_records']):
            flags.append('oil_input_requires_review')
        if source == 'M-1' and 'AC-13-3' in grade:
            flags.append('product_alias_unconfirmed_90_suffix')
        flags += ['test_conditions_unverified', 'rap_mass_basis_unconfirmed']
        conditional = [a for a in ['B/B-2 AC-10' if source in ['B','B-2'] else '',
                                   '1G AC-13/AC-13-1' if source == '1G' else '',
                                   '1K AC-13/AC-10' if source == '1K' else '',
                                   'M-1 90#/90#-3' if source == 'M-1' else ''] if a]
        session.add(Recipe(id=recipe_id, dataset_id=dsid, source=source, gradation=grade, product=product,
            oil_stone_ratio_pct=oil, rap_nominal_pct=gradation['rap_nominal_pct'] if gradation else None,
            payload={'fixture_id': fixture_index.get(key), 'gradation': gradation, 'flags': sorted(set(flags)),
                     'identity_links': {'gradation': {'state':'exact' if gradation else 'unresolved',
                         'key':[source,grade], 'reference':f'级配!B{gradation["start_row"]}' if gradation else None},
                         'aliases':[{'state':'conditional','alias':a,'applied':False} for a in conditional]},
                     'specimen_records': metric_items['_records'], 'evidence_kind': 'historical_record_aggregate'}))
        for metric, items in metric_items.items():
            if metric == '_records' or not items:
                continue
            suspected_mean = metric == 'DS' and len(items) == 3 and abs(items[2]['value'] - (items[0]['value']+items[1]['value'])/2) < 1e-7
            if suspected_mean:
                items[2]['suspected_derived_mean'] = True
                suspected.append(recipe_id)
            independent = items[:2] if suspected_mean else items
            vals = [x['value'] for x in independent]
            session.add(Measurement(id=f'{recipe_id}:{metric}', dataset_id=dsid, recipe_id=recipe_id,
                metric=metric, value=sum(vals)/len(vals), unit=METRICS[metric]['unit'], payload={
                    'raw_labels': items, 'range': [min(vals), max(vals)], 'n_recorded': len(items),
                    'n_independent_assumed': None if metric in ['TSR','residual'] else len(independent), 'suspected_derived_mean': suspected_mean,
                    'aggregation': 'stored_group_ratio' if metric in ['TSR','residual'] else 'mean_of_records',
                    'conditions_status': 'unverified', 'source_kind': 'official_workbook',
                    'paired_group_rows': [r for r in metric_items['_records'] if r['sheet'] == items[0]['sheet']]}))
    ds.summary = {'layout_version': 'workbook-v1', 'sheets': sheets, 'specimen_counts': counts,
        'specimen_total': sum(counts.values()), 'gradation_blocks': len(blocks),
        'source_codes': len({k[0] for k in blocks}), 'operational_recipes': len(groups),
        'formula_errors': errors, 'curve_mismatches': mismatches,
        'suspected_mean_DS_groups': len(suspected), 'unmatched_gradation_keys': no_grade,
        'scope': '实验室混合料数据；试件行数不是独立配方数；来源代码不代表已确认独立料源',
        'conditional_aliases': ['B/B-2 AC-10', '1G AC-13/AC-13-1', '1K AC-13/AC-10', 'M-1 90#/90#-3'],
        'identity_states': ['exact','confirmed_alias','conditional','unresolved'],
        'confirmed_aliases':[], 'metric_coverage':{m:{'operational_groups':sum(bool(g.get(m)) for g in groups.values()),
            'stored_numeric_labels':sum(len(g.get(m,[])) for g in groups.values())} for m in METRICS}}
    session.add(ds)
    session.commit()
    book.close()
    formulas.close()
    return ds, True


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook', type=Path)
    args = parser.parse_args()
    initialize(engine)
    with Session(engine) as session:
        dataset, created = import_workbook(args.workbook, session)
        print(json.dumps({'id': dataset.id, 'created': created, 'sha256': dataset.sha256,
                          'summary': dataset.summary}, ensure_ascii=False, indent=2))
