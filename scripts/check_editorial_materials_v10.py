"""Verify that source paragraphs, structured values and final documents agree."""
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from lxml import etree as E
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
ART=ROOT/'artifacts/editorial-v10'
OUT=ROOT/'output/submission-v10'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s):return re.sub(r'\s+','',s)

def main():
    paper=OUT/'07_作品论文/作品论文.docx';doc=Document(paper)
    body=[norm(p.text) for p in doc.paragraphs]
    checked=0;missing=[]
    for p in sorted((ROOT/'docs/paper-v10').glob('*.md')):
        for block in p.read_text(encoding='utf-8').strip().split('\n\n'):
            if block.startswith('['):continue
            value=norm(re.sub(r'^#+ ','',block))
            checked+=1
            if value not in body:missing.append({'file':p.name,'text':value[:90]})
    assert not missing,missing
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with ZipFile(paper) as z:x=E.fromstring(z.read('word/document.xml'))
    counts={'native_tables':len(x.xpath('//w:tbl',namespaces=ns)),
            'native_equations':len(x.xpath('//m:oMath',namespaces=ns)),
            'inline_figures':len(doc.inline_shapes)}
    assert counts=={'native_tables':30,'native_equations':7,'inline_figures':12}
    references=json.loads((ROOT/'docs/paper-v10/references.json').read_text(encoding='utf-8'))
    assert len(references)==12 and len({r['url'] for r in references})==12
    form=json.loads((ROOT/'docs/project-summary-v10.json').read_text(encoding='utf-8'))
    with ZipFile(OUT/'01_必交文档/项目简表.docx') as z:summary=E.fromstring(z.read('word/document.xml'))
    rows=summary.find('w:body/w:tbl',ns).findall('w:tr',ns)
    for index,value in form['sections'].items():
        cell=rows[int(index)].findall('w:tc',ns)[1]
        assert ''.join(cell.xpath('.//w:t/text()',namespaces=ns))==value
    for source,relative in [('docs/submission-final-v10.md','01_必交文档/项目说明书.docx'),('docs/business-plan-v10.md','04_商业计划书/商业计划书.docx')]:
        paragraphs={norm(p.text) for p in Document(OUT/relative).paragraphs}
        for block in (ROOT/source).read_text(encoding='utf-8').strip().split('\n\n'):
            if block.startswith('#'):continue
            assert norm(block) in paragraphs,(source,block[:50])
    pages={p.relative_to(OUT).as_posix():len(PdfReader(p).pages) for p in OUT.rglob('*.pdf')}
    assert pages['07_作品论文/作品论文.pdf']==50
    assert sha(ROOT/'output/submission/匿名提交包.zip')=='8dc932e9452e4df4ccba13aba85e12573d0b362af776f62fad6c2e49ddec2eec'
    assert sha(ROOT/'data/workbench.sqlite')=='624862ee54e8282eb00a81713f0bc2c7d800fa5ddd435a5e1089b395278464a6'
    record={'created_on':'2026-10-24','status':'passed','paper_source_blocks_checked':checked,
            'paper_source_blocks_missing':0,**counts,'references':len(references),'pdf_page_counts':pages,
            'summary_sections_match':4,'official_body_paragraphs_match':True,
            'frozen_source_archive_unchanged':True,'source_database_unchanged':True}
    (ART/'content-verification.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    build=json.loads((ART/'paper-build.json').read_text(encoding='utf-8'))
    build.update(sha256=sha(paper),pdf_sha256=sha(OUT/'07_作品论文/作品论文.pdf'),pages=50,render_status='rendered_awaiting_final_visual_record')
    (ART/'paper-build.json').write_text(json.dumps(build,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    reading='\n\n'.join(p.read_text(encoding='utf-8').strip() for p in sorted((ROOT/'docs/paper-v10').glob('*.md')))
    (ART/'manuscript-reading-copy.txt').write_text(reading+'\n',encoding='utf-8')
    print(json.dumps(record,ensure_ascii=False))

if __name__=='__main__':main()
