"""Build the long-form competition paper with native tables and Word equations."""
import argparse
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from paper_tables_v10 import catalogue
from paper_figures_v10 import build_figures

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'docs/paper-v10'
ART=ROOT/'artifacts/editorial-v10'


def math_run(value):
    node=OxmlElement('m:r')
    props=OxmlElement('m:rPr')
    style=OxmlElement('m:sty');style.set(qn('m:val'),'p');props.append(style);node.append(props)
    text=OxmlElement('m:t');text.text=value;node.append(text)
    return node


def script(base, index, kind='sub'):
    node=OxmlElement('m:sSub' if kind=='sub' else 'm:sSup')
    e=OxmlElement('m:e');e.append(math_run(base));node.append(e)
    element=OxmlElement('m:sub' if kind=='sub' else 'm:sup')
    element.append(math_run(index));node.append(element)
    return node


def fraction(top, bottom):
    node=OxmlElement('m:f')
    for key,value in [('num',top),('den',bottom)]:
        el=OxmlElement('m:'+key)
        for part in value: el.append(math_run(part) if isinstance(part,str) else part)
        node.append(el)
    return node


def sigma(lower, upper, parts):
    n=OxmlElement('m:nary');props=OxmlElement('m:naryPr')
    chr=OxmlElement('m:chr');chr.set(qn('m:val'),'∑');props.append(chr)
    loc=OxmlElement('m:limLoc');loc.set(qn('m:val'),'subSup');props.append(loc)
    if not upper:
        hide=OxmlElement('m:supHide');hide.set(qn('m:val'),'1');props.append(hide)
    n.append(props)
    for key,value in [('sub',lower),('sup',upper)]:
        el=OxmlElement('m:'+key);el.append(math_run(value));n.append(el)
    e=OxmlElement('m:e')
    for part in parts: e.append(math_run(part) if isinstance(part,str) else part)
    n.append(e)
    return n


def equation_parts(key):
    return {
        'feasible':['F(P,d) = { r ∈ P ∣ ∀t ∈ T(d), c(r,t) = 1 }'],
        'dominance':['A ≻ B ⇔ [∀i, ',script('y','i'),'(A) ≥ ',script('y','i'),'(B)] ∧ [∃j, ',
                     script('y','j'),'(A) > ',script('y','j'),'(B)]'],
        'conflict':['∀r ∈ P, ',script('F','r'),' ∩ U ≠ ∅'],
        'minimality':['∀t ∈ U, ∃r ∈ P: ',script('F','r'),' ∩ (U ∖ {t}) = ∅'],
        'mae':['MAE = ',fraction(['1'],['n']),sigma('i=1','n',['|',script('ŷ','i'),' − ',script('y','i'),'|'])],
        'cost':['C(s) = ',sigma('k','',[script('q','k'),script('p','k')]),' + ',
                sigma('t','',[fraction([script('C','s,t')],[script('(1+r)','t','sup')])])],
        'carbon':['E(s) = ',sigma('a','',[script('A','a'),'(s)',script('f','a')])],
    }[key]


def apply_font(style, size, east='宋体', latin='Times New Roman', bold=False):
    style.font.name=latin;style.font.size=Pt(size);style.font.bold=bold;style.font.italic=False
    style.font.color.rgb=RGBColor(0,0,0)
    props=style.element.get_or_add_rPr()
    fonts=props.find(qn('w:rFonts'))
    if fonts is None: fonts=OxmlElement('w:rFonts');props.insert(0,fonts)
    for script_name,value in [('ascii',latin),('hAnsi',latin),('eastAsia',east),('cs',latin)]:
        fonts.set(qn('w:'+script_name),value)
    for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        fonts.attrib.pop(qn('w:'+key),None)
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0')
    style.element.get_or_add_pPr().append(snap)


def field(paragraph,instruction):
    run=paragraph.add_run()
    for typ in ['begin','separate','end']:
        element=OxmlElement('w:fldChar');element.set(qn('w:fldCharType'),typ)
        if typ=='begin':
            run._r.append(element)
            text=OxmlElement('w:instrText');text.set(qn('xml:space'),'preserve');text.text=instruction
            run._r.append(text)
        elif typ=='separate':
            run._r.append(element)
            value=OxmlElement('w:t');value.text='';run._r.append(value)
        else:run._r.append(element)


def build(output,line_spacing):
    data=json.loads((ART/'evidence.json').read_text(encoding='utf-8'))
    tables=catalogue(data)
    figures=build_figures(data,ART/'figures')
    doc=Document()
    section=doc.sections[0]
    # A separate technical attachment; official forms keep their retained A4 geometry.
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.left_margin=section.right_margin=Inches(1)
    section.top_margin=Inches(.92);section.bottom_margin=Inches(.9)
    section.header_distance=Inches(.4);section.footer_distance=Inches(.42)
    section.different_first_page_header_footer=True
    for name,size,east,bold in [
        ('Normal',12,'宋体',False),('Title',26,'黑体',True),('Subtitle',15,'宋体',False),
        ('Heading 1',18,'黑体',True),('Heading 2',14,'黑体',True),
        ('Heading 3',12,'黑体',True),('Caption',10.5,'宋体',False),
        ('Header',9,'宋体',False),('Footer',9,'宋体',False)]:
        apply_font(doc.styles[name],size,east,bold=bold)
    normal=doc.styles['Normal'].paragraph_format
    normal.line_spacing=Pt(line_spacing);normal.line_spacing_rule=WD_LINE_SPACING.EXACTLY
    normal.first_line_indent=Pt(24);normal.space_after=Pt(4);normal.widow_control=True
    # Use plain headings and a compact, readable contents list.
    for style in doc.styles:
        for node in list(style.element.iter(qn('w:pBdr'))):
            node.getparent().remove(node)
    for level in [1,2]:
        name=f'TOC {level}'
        style=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,1)
        apply_font(style,10.5)
        fmt=style.paragraph_format
        fmt.line_spacing=Pt(15);fmt.space_after=Pt(2);fmt.space_before=Pt(0)
        fmt.first_line_indent=Pt(0);fmt.left_indent=Pt(13 if level==2 else 0)
    for name in ['Heading 1','Heading 2','Heading 3']:
        p=doc.styles[name].paragraph_format
        p.first_line_indent=Pt(0);p.space_before=Pt(13);p.space_after=Pt(7)
        p.keep_with_next=True;p.keep_together=True;p.line_spacing=1.2
    doc.styles['Heading 1'].paragraph_format.page_break_before=True
    cp=doc.styles['Caption'].paragraph_format
    cp.first_line_indent=Pt(0);cp.line_spacing=1.2;cp.space_before=Pt(7);cp.space_after=Pt(5)
    if 'Hyperlink' in doc.styles:
        doc.styles['Hyperlink'].font.color.rgb=RGBColor(0,0,0)
        doc.styles['Hyperlink'].font.underline=False
    header=section.header.paragraphs[0]
    header.alignment=WD_ALIGN_PARAGRAPH.CENTER;header.paragraph_format.first_line_indent=Pt(0)
    header.add_run('数智选材 道路养护材料推荐与协同设计')
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Pt(0)
    field(footer,' PAGE ')
    for label,size,top in [('数智选材',30,90),('道路养护材料推荐与协同设计',22,24),('作品论文',16,42)]:
        p=doc.add_paragraph(style='Title' if size>=22 else 'Subtitle')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.space_before=Pt(top);p.paragraph_format.space_after=Pt(10)
        cover_run=p.add_run(label);cover_run.font.size=Pt(size);cover_run.font.italic=False
    p=doc.add_paragraph('企业揭榜赛道 HY-01 课题一',style='Subtitle')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    for run in p.runs:run.font.italic=False
    p=doc.add_paragraph('2026年10月',style='Subtitle')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    for run in p.runs:run.font.italic=False
    chapter='0';table_no=figure_no=eq_no=0
    used_tables=[];used_figures=[];used_equations=[]
    def body(value):
        p=doc.add_paragraph(value)
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        return p
    def caption(value,keep=True):
        p=doc.add_paragraph(value,style='Caption')
        p.paragraph_format.keep_with_next=keep
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        return p
    def source_note(value):
        p=doc.add_paragraph()
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.15
        p.paragraph_format.space_after=Pt(9);p.paragraph_format.space_before=Pt(3)
        p.add_run(value).font.size=Pt(9)
    def add_table(key):
        nonlocal table_no
        table_no+=1; info=tables[key]
        caption(f'表{chapter}-{table_no} {info["title"]}')
        t=doc.add_table(rows=1,cols=len(info['headers']))
        t.autofit=False
        widths=info['widths'] or [165/len(info['headers'])]*len(info['headers'])
        for col,width in zip(t.columns,widths):col.width=Mm(width)
        values=[info['headers']]+info['rows']
        for index,vals in enumerate(values):
            row=t.rows[0] if index==0 else t.add_row()
            props=row._tr.get_or_add_trPr()
            cant=OxmlElement('w:cantSplit');props.append(cant)
            if index==0:
                repeat=OxmlElement('w:tblHeader');props.append(repeat)
            for j,(cell,value) in enumerate(zip(row.cells,vals)):
                cell.width=Mm(widths[j])
                cell.vertical_alignment=1
                tcpr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    b=OxmlElement('w:'+side);b.set(qn('w:val'),'single')
                    b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
                tcpr.append(borders)
                margin=OxmlElement('w:tcMar')
                for side,v in [('top',80),('bottom',80),('left',90),('right',90)]:
                    m=OxmlElement('w:'+side);m.set(qn('w:w'),str(v));m.set(qn('w:type'),'dxa');margin.append(m)
                tcpr.append(margin)
                if index==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F2F1');tcpr.append(shade)
                p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.line_spacing=1.18;p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.keep_with_next=len(values)<=5 and index<len(values)-1
                run=p.add_run(str(value));run.font.size=Pt(10.5);run.bold=index==0
        source_note(info['source']);used_tables.append(key)
    def add_figure(key):
        nonlocal figure_no
        figure_no+=1;info=figures[key]
        p=doc.add_paragraph()
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
        p.paragraph_format.line_spacing=1.0
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(info['path'],width=Inches(6.3))
        caption(f'图{chapter}-{figure_no} {info["caption"]}',keep=True)
        source_note(info['source']);used_figures.append(key)
    def add_equation(key):
        nonlocal eq_no
        eq_no+=1
        p=doc.add_paragraph()
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(7)
        p.paragraph_format.line_spacing=1.2
        p.paragraph_format.space_after=Pt(9);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        mpara=OxmlElement('m:oMathPara');m=OxmlElement('m:oMath')
        for part in equation_parts(key):m.append(math_run(part) if isinstance(part,str) else part)
        m.append(math_run(f'     ({chapter}-{eq_no})'))
        mpara.append(m);p._p.append(mpara);used_equations.append(key)
    def parse(path):
        nonlocal chapter,table_no,figure_no,eq_no
        for block in path.read_text(encoding='utf-8').strip().split('\n\n'):
            if block.startswith('# '):
                value=block[2:]
                match=re.match(r'第(\d+)章',value)
                appendix=re.match(r'附录([A-Z])',value)
                chapter=match.group(1) if match else appendix.group(1) if appendix else '0'
                table_no=figure_no=eq_no=0
                doc.add_heading(value,level=1)
            elif block.startswith('## '):
                heading=doc.add_heading(block[3:],level=2)
                if block.startswith('## A2 '):heading.paragraph_format.page_break_before=True
            elif block.startswith('[TABLE:'):add_table(block[7:-1])
            elif block.startswith('[FIGURE:'):add_figure(block[8:-1])
            elif block.startswith('[EQUATION:'):add_equation(block[10:-1])
            elif block=='[CODE:startup]':
                for value in ['powershell -ExecutionPolicy Bypass -File scripts/start-creative-demo.ps1',
                              '.venv\\Scripts\\python.exe -X utf8 scripts/replay_creative_workbench_v7.py']:
                    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
                    p.paragraph_format.line_spacing=1.15
                    r=p.add_run(value);r.font.name='Consolas';r.font.size=Pt(9.5)
            else:body(block.replace('\n',' '))
    parse(SOURCE/'00-abstract.md')
    p=doc.add_paragraph('目录',style='Title')
    p.paragraph_format.page_break_before=True;p.paragraph_format.first_line_indent=Pt(0)
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
    field(p,r' TOC \o "1-2" \h \z \u ')
    for path in sorted(SOURCE.glob('*.md')):
        if re.match(r'(0[1-9]|1[0-2])-',path.name):parse(path)
    doc.add_heading('参考文献',level=1)
    refs=json.loads((SOURCE/'references.json').read_text(encoding='utf-8'))
    for ref in refs:
        p=doc.add_paragraph(f'[{ref["id"]}] {ref["text"]}\n{ref["url"]}')
        p.paragraph_format.first_line_indent=Pt(-22);p.paragraph_format.left_indent=Pt(22)
        p.paragraph_format.line_spacing=Pt(17);p.paragraph_format.space_after=Pt(9)
        p.paragraph_format.keep_together=True
        for run in p.runs:run.font.size=Pt(10.5)
    parse(SOURCE/'13-appendices.md')
    settings=doc.settings.element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    doc.core_properties.title='数智选材 道路养护材料推荐与协同设计'
    doc.core_properties.subject='作品论文'
    doc.core_properties.author=''
    doc.core_properties.last_modified_by=''
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    manuscript='\n\n'.join(p.read_text(encoding='utf-8').strip() for p in sorted(SOURCE.glob('*.md')))
    (ART/'manuscript-reading-copy.txt').write_text(manuscript+'\n',encoding='utf-8')
    result={'output':output.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(output.read_bytes()).hexdigest(),
            'body_chinese_characters':len(re.findall('[\u4e00-\u9fff]',manuscript)),
            'native_tables':len(used_tables),'figures':len(used_figures),'native_equations':len(used_equations),
            'table_keys':used_tables,'figure_keys':used_figures,'equation_keys':used_equations,
            'references':len(refs),'body_font_pt':12,'line_spacing_pt':line_spacing,
            'page_size':'Letter portrait','render_status':'pending'}
    (ART/'paper-build.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',type=Path,default=ROOT/'output/submission-v10/07_作品论文/作品论文.docx')
    parser.add_argument('--line-spacing',type=float,default=18)
    args=parser.parse_args()
    build(args.output.resolve(),args.line_spacing)
